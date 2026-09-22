"""Standalone parser subprocess. Do not import the application or model SDK here."""

import json
import math
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree


class InvalidDocument(Exception):
    def __init__(self, code: str, message: str):
        self.code, self.message = code, message


def resource_limits(limits: dict) -> None:
    import resource

    memory = limits["memory_mb"] * 1024 * 1024
    cpu = math.ceil(limits["cpu_seconds"])
    resource.setrlimit(resource.RLIMIT_CPU, (cpu, cpu + 1))
    resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    # macOS rejects practical address-space limits for modern Python. Its parent
    # enforces an RSS budget instead; Linux deployment additionally uses RLIMIT_AS.
    if sys.platform != "darwin":
        resource.setrlimit(resource.RLIMIT_AS, (memory, memory))


def restrict_operations(event: str, args: tuple) -> None:
    if (
        event.startswith(("socket.", "subprocess.", "os.exec", "os.spawn"))
        or event == "os.system"
    ):
        raise PermissionError(
            "Network and process execution are disabled in document parsers"
        )
    if event == "open":
        mode, flags = args[1], args[2]
        if (
            isinstance(mode, str) and any(flag in mode for flag in "wax+")
        ) or flags & 0x3:
            raise PermissionError("Filesystem writes are disabled in document parsers")


class Collector:
    def __init__(self, limits: dict):
        self.limits = limits
        self.characters = 0
        self.blocks = []

    def add(self, text: str, **metadata) -> None:
        text = text.strip()
        if not text:
            return
        self.characters += len(text)
        if self.characters > self.limits["max_characters"] or len(self.blocks) >= 10000:
            raise InvalidDocument(
                "EXTRACTED_TEXT_TOO_LARGE",
                "Extracted document text exceeds the configured limit.",
            )
        self.blocks.append({"text": text, "metadata": metadata})


def parse_text(path: Path, output: Collector, markdown: bool) -> None:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeError as exc:
        raise InvalidDocument(
            "INVALID_TEXT_ENCODING", "TXT and Markdown files must use UTF-8 encoding."
        ) from exc
    if "\x00" in text:
        raise InvalidDocument("FILE_TYPE_MISMATCH", "The file contains binary data.")
    section = []
    buffer = []
    start = 1
    fenced = False
    for line_number, line in enumerate(text.splitlines(), 1):
        if markdown and line.lstrip().startswith(("```", "~~~")):
            fenced = not fenced
        heading = (
            re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
            if markdown and not fenced
            else None
        )
        if heading:
            output.add(
                "\n".join(buffer),
                section=" / ".join(section),
                line_start=start,
                line_end=line_number - 1,
            )
            level = len(heading[1])
            section = section[: level - 1] + [heading[2]]
            buffer = [line]
            start = line_number
        else:
            buffer.append(line)
    output.add(
        "\n".join(buffer),
        section=" / ".join(section),
        line_start=start,
        line_end=len(text.splitlines()),
    )


def inspect_docx(path: Path, limits: dict) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > limits["max_entries"]:
                raise InvalidDocument(
                    "DOCX_ARCHIVE_LIMIT", "The DOCX archive contains too many entries."
                )
            total = 0
            names = set()
            for item in entries:
                if item.filename in names:
                    raise InvalidDocument(
                        "INVALID_DOCX", "The DOCX archive contains duplicate entries."
                    )
                names.add(item.filename)
                if (
                    "\\" in item.filename
                    or item.filename.startswith("/")
                    or ".." in Path(item.filename).parts
                ):
                    raise InvalidDocument(
                        "INVALID_DOCX", "The DOCX archive contains an unsafe entry."
                    )
                total += item.file_size
                if (
                    total > limits["max_uncompressed_bytes"]
                    or item.file_size
                    > max(1, item.compress_size) * limits["max_compression_ratio"]
                ):
                    raise InvalidDocument(
                        "DOCX_ARCHIVE_LIMIT",
                        "The DOCX archive exceeds decompression limits.",
                    )
                if item.flag_bits & 1:
                    raise InvalidDocument(
                        "ENCRYPTED_DOCUMENT", "Encrypted DOCX files are not supported."
                    )
                if "vbaproject" in item.filename.lower():
                    raise InvalidDocument(
                        "UNSUPPORTED_FILE_TYPE",
                        "Macro-enabled documents are not supported.",
                    )
            if not {"[Content_Types].xml", "word/document.xml"}.issubset(names):
                raise InvalidDocument(
                    "INVALID_DOCX", "The archive is not a DOCX document."
                )
            # Check limits before materializing XML; no archive entry is extracted.
            document_xml = archive.read("word/document.xml")
            if (
                b"<!DOCTYPE" in document_xml.upper()
                or b"<!ENTITY" in document_xml.upper()
            ):
                raise InvalidDocument(
                    "INVALID_DOCX", "DTD and entity declarations are not allowed."
                )
            root = ElementTree.fromstring(document_xml)
            namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            explicit_breaks = sum(
                element.attrib.get(namespace + "type") == "page"
                for element in root.iter(namespace + "br")
            )
            if explicit_breaks + 1 > limits["max_pages"]:
                raise InvalidDocument(
                    "PAGE_LIMIT", "The document exceeds the explicit page-break limit."
                )
    except zipfile.BadZipFile as exc:
        raise InvalidDocument("INVALID_DOCX", "The DOCX archive is corrupt.") from exc


def parse_docx(path: Path, output: Collector) -> None:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    inspect_docx(path, output.limits)
    document = Document(path)
    sections = []
    buffer = []
    paragraph_index, table_index, start = 0, 0, 1

    def flush(end: int | None = None) -> None:
        output.add(
            "\n\n".join(buffer),
            section=" / ".join(sections),
            paragraph_start=start,
            paragraph_end=paragraph_index if end is None else end,
        )
        buffer.clear()

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_index += 1
            text = block.text.strip()
            style = block.style.name if block.style is not None else ""
            if style.startswith("Heading ") and style[8:].isdigit():
                flush(paragraph_index - 1)
                level = max(1, int(style[8:]))
                sections = sections[: level - 1] + [text]
                start = paragraph_index
            buffer.append(text)
        elif isinstance(block, Table):
            flush()
            table_index += 1
            rows = [
                [cell.text.strip().replace("\n", " ") for cell in row.cells]
                for row in block.rows
            ]
            if not rows:
                continue
            header = rows[0]
            for row_index, row in enumerate(
                rows[1:] or rows, 2 if len(rows) > 1 else 1
            ):
                # Every row keeps its header when a long table spans chunks.
                output.add(
                    " | ".join(header) + "\n" + " | ".join(row),
                    section=" / ".join(sections),
                    table=table_index,
                    row=row_index,
                    table_headers=header,
                )
            start = paragraph_index + 1
    flush()


def parse_pdf(path: Path, output: Collector) -> tuple[int, list[str]]:
    from pypdf import PdfReader

    document = PdfReader(path, strict=False)
    if document.is_encrypted:
        raise InvalidDocument(
            "ENCRYPTED_DOCUMENT", "Encrypted PDF files are not supported."
        )
    count = len(document.pages)
    if count > output.limits["max_pages"]:
        raise InvalidDocument(
            "PAGE_LIMIT", "The PDF exceeds the configured page limit."
        )
    blank_pages = 0
    for page_index, page in enumerate(document.pages, 1):
        text = page.extract_text() or ""
        if not text.strip():
            blank_pages += 1
            continue
        output.add(text, page=page_index, page_label=str(page_index))
    warnings = (
        ["Some PDF pages have no extractable text; OCR is not enabled."]
        if blank_pages
        else []
    )
    return count, warnings


def main() -> None:
    path = Path(sys.argv[1])
    limits = json.loads(sys.argv[2])
    resource_limits(limits)
    sys.addaudithook(restrict_operations)
    output = Collector(limits)
    page_count, warnings = None, []
    try:
        extension = limits["extension"]
        if extension in {".txt", ".md"}:
            parse_text(path, output, extension == ".md")
        elif extension == ".pdf":
            page_count, warnings = parse_pdf(path, output)
        elif extension == ".docx":
            parse_docx(path, output)
        if not output.blocks:
            code = "OCR_REQUIRED" if extension == ".pdf" else "NO_TEXT"
            raise InvalidDocument(
                code, "No text could be extracted; scanned PDF files require OCR."
            )
        result = {
            "blocks": output.blocks,
            "page_count": page_count,
            "warnings": warnings,
        }
    except InvalidDocument as exc:
        result = {"error": {"code": exc.code, "message": exc.message}}
    except MemoryError:
        result = {
            "error": {
                "code": "PARSE_RESOURCE_LIMIT",
                "message": "Document parsing exceeded the memory limit.",
            }
        }
    except Exception:  # noqa: BLE001 -- sanitize all third-party parser errors.
        # Parser errors can contain original text, paths or XML. Do not leak them.
        result = {
            "error": {
                "code": "INVALID_DOCUMENT",
                "message": "The document could not be parsed.",
            }
        }
    sys.stdout.write(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
