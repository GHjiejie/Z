"""Offline ingestion tests exercise actual document readers and LlamaIndex."""

import hashlib
import io
import tempfile
import threading
import time
import unittest
import zipfile
from dataclasses import asdict
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import httpx
from docx import Document as WordDocument
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, StreamObject

from llamaindex_service.config import Settings
from llamaindex_service.contracts import ServiceError
from llamaindex_service.ingestion import IngestionProcessor, parse_document
from llamaindex_service.retrieval.embeddings import EmbeddingError
from llamaindex_service.storage import LocalStorage
from llamaindex_service.workers import Worker
from llamaindex_service.workers.__main__ import require_worker_readiness
from llamaindex_service.workers.runner import classify_failure


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name)
        self.settings = Settings(
            _env_file=None,
            environment="test",
            model_mode="mock",
            embedding_dimension=64,
            storage_path=self.path,
            chunk_size=64,
            chunk_overlap=8,
        )
        self.storage = LocalStorage(self.path)

    def save(self, filename, content):
        path = self.path / filename
        path.write_bytes(content)
        return path

    def test_markdown_sections_and_integrity(self):
        text = b"# Main\nProduct onboarding.\n## Security\nUse signed tokens."
        path = self.save("guide.md", text)
        result = parse_document(
            path, path.name, self.settings, hashlib.sha256(text).hexdigest()
        )
        self.assertEqual(
            [block.metadata["section"] for block in result.blocks],
            ["Main", "Main / Security"],
        )
        self.assertIsNone(result.page_count)
        self.assertEqual(result.blocks[1].metadata["line_start"], 3)
        with self.assertRaises(ServiceError) as raised:
            parse_document(path, path.name, self.settings, "0" * 64)
        self.assertEqual(raised.exception.code, "CONTENT_HASH_MISMATCH")

    def test_encoding_extracted_size_and_preexisting_cancel(self):
        path = self.save("invalid.txt", b"\xff\xfeabc")
        with self.assertRaises(ServiceError) as raised:
            parse_document(path, path.name, self.settings)
        self.assertEqual(raised.exception.code, "INVALID_TEXT_ENCODING")
        path = self.save("large.txt", b"too much text")
        settings = self.settings.model_copy(update={"max_extracted_characters": 3})
        with self.assertRaises(ServiceError) as raised:
            parse_document(path, path.name, settings)
        self.assertEqual(raised.exception.code, "EXTRACTED_TEXT_TOO_LARGE")
        cancellation = threading.Event()
        cancellation.set()
        with self.assertRaises(ServiceError) as raised:
            parse_document(path, path.name, self.settings, cancelled=cancellation)
        self.assertEqual(raised.exception.code, "JOB_CANCELLED")

    def test_docx_paragraphs_tables_and_no_invented_pages(self):
        document = WordDocument()
        document.add_heading("Enterprise onboarding", level=1)
        document.add_paragraph("Employees must register a security token.")
        table = document.add_table(rows=3, cols=2)
        for cells, values in zip(
            table.rows,
            [["Product", "Region"], ["DEV-2026", "Shanghai"], ["EDGE-2026", "Beijing"]],
            strict=True,
        ):
            for cell, value in zip(cells.cells, values, strict=True):
                cell.text = value
        document.add_heading("Support", level=2)
        document.add_paragraph("The service desk accepts weekday tickets.")
        path = self.path / "handbook.docx"
        document.save(path)
        result = parse_document(path, path.name, self.settings)
        self.assertIsNone(result.page_count)
        self.assertTrue(all("page" not in block.metadata for block in result.blocks))
        table_blocks = [block for block in result.blocks if "table" in block.metadata]
        self.assertEqual(len(table_blocks), 2)
        self.assertTrue(
            all(block.text.startswith("Product | Region") for block in table_blocks)
        )
        self.assertEqual(
            result.blocks[-1].metadata["section"], "Enterprise onboarding / Support"
        )

    def test_archive_bomb_is_rejected_before_decompression(self):
        path = self.path / "bomb.docx"
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "types")
            archive.writestr("word/document.xml", "x" * 200000)
        with self.assertRaises(ServiceError) as raised:
            parse_document(path, path.name, self.settings)
        self.assertEqual(raised.exception.code, "DOCX_ARCHIVE_LIMIT")

    def test_parser_wall_timeout(self):
        path = self.save("manual.txt", b"A short document.")
        settings = self.settings.model_copy(update={"parse_timeout_seconds": 0.00001})
        with self.assertRaises(ServiceError) as raised:
            parse_document(path, path.name, settings)
        self.assertEqual(raised.exception.code, "PARSE_TIMEOUT")

    def test_pdf_page_source_ocr_and_page_limit(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = StreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 250 Td (Enterprise knowledge service) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.add_blank_page(width=300, height=300)
        path = self.path / "report.pdf"
        writer.write(path)
        result = parse_document(path, path.name, self.settings)
        self.assertEqual(result.page_count, 2)
        self.assertEqual(result.blocks[0].metadata["page"], 1)
        self.assertIn("Enterprise knowledge", result.blocks[0].text)
        self.assertTrue(result.warnings)
        settings = self.settings.model_copy(update={"max_pages": 1})
        with self.assertRaises(ServiceError) as raised:
            parse_document(path, path.name, settings)
        self.assertEqual(raised.exception.code, "PAGE_LIMIT")
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        writer.write(path)
        with self.assertRaises(ServiceError) as raised:
            parse_document(path, path.name, self.settings)
        self.assertEqual(raised.exception.code, "OCR_REQUIRED")

    def test_real_pipeline_is_deterministic_and_version_scoped(self):
        text = "# Access\n企业员工必须使用安全令牌。文档知识服务支持检索和问答。\n" * 12
        stored = self.storage.write(io.BytesIO(text.encode()), "guide.md")
        job = {
            **asdict(stored),
            "tenant_id": "t1",
            "project_id": "p1",
            "kb_id": "kb",
            "document_id": "document",
            "version_id": "v1",
            "generation_id": "gen1",
        }
        processor = IngestionProcessor(self.settings, self.storage)
        phases = []
        first = processor.process(job, lambda phase, count: phases.append(phase))
        second = processor.process(job)
        self.assertEqual(first, second)
        self.assertGreater(len(first), 1)
        self.assertTrue(all(len(chunk["embedding"]) == 64 for chunk in first))
        self.assertTrue(
            all(chunk["metadata"]["section"] == "Access" for chunk in first)
        )
        self.assertTrue(all(chunk["lexical_text"] for chunk in first))
        self.assertIn("EMBEDDING", phases)
        third = processor.process({**job, "version_id": "v2"})
        self.assertTrue(
            {chunk["id"] for chunk in first}.isdisjoint(chunk["id"] for chunk in third)
        )
        with self.assertRaises(ServiceError) as raised:
            processor.process({**job, "embedding_dimension": 128})
        self.assertEqual(raised.exception.code, "INDEX_MODEL_MISMATCH")


class FakeRepository:
    def __init__(self, kind="ingest", valid=True):
        self.job = {"job_id": "job", "claim_token": 1, "kind": kind}
        self.valid = valid
        self.published = None
        self.failed = None
        self.cleaned = False
        self.phases = []
        self.keys = []
        self.heartbeats = 0

    def claim_job(self, worker_id, lease_seconds):
        job, self.job = self.job, None
        return job

    def heartbeat(self, *args):
        self.heartbeats += 1
        return self.valid

    def job_phase(self, job_id, token, phase, completed=0):
        self.phases.append(phase)
        return self.valid

    def publish_chunks(self, job_id, token, chunks, **kwargs):
        if not self.valid:
            raise ServiceError("JOB_CANCELLED", "Claim invalid", 409)
        self.published = chunks

    def fail_job(self, job_id, token, code, message, **kwargs):
        self.failed = {"code": code, "message": message, **kwargs}

    def cleanup_objects(self, *args):
        return self.keys

    def complete_cleanup(self, *args):
        self.cleaned = True


class FakeProcessor:
    def __init__(self, failure=None, delay=0):
        self.failure, self.delay = failure, delay

    def process(self, job, phase, cancelled):
        phase("PARSING", 0)
        time.sleep(self.delay)
        if self.failure:
            raise self.failure
        phase("EMBEDDING", 1)
        return [{"id": "chunk"}]


class WorkerTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.settings = Settings(
            _env_file=None,
            environment="test",
            model_mode="mock",
            storage_path=self.directory.name,
        )
        self.storage = LocalStorage(self.directory.name)

    def test_publish_and_cancel_fencing(self):
        repo = FakeRepository()
        worker = Worker(self.settings, repo, self.storage, FakeProcessor())
        self.assertTrue(worker.run_once())
        self.assertEqual(repo.published, [{"id": "chunk"}])
        self.assertFalse(worker.run_once())
        repo = FakeRepository(valid=False)
        self.assertTrue(
            Worker(self.settings, repo, self.storage, FakeProcessor()).run_once()
        )
        self.assertIsNone(repo.published)
        self.assertIsNone(repo.failed)

    def test_production_startup_requires_worker_role_and_fails_closed(self):
        settings = SimpleNamespace(environment="production")
        repository = Mock()
        repository.health.return_value = True
        require_worker_readiness(settings, repository)
        repository.health.assert_called_once_with(worker=True)
        repository.health.return_value = False
        with self.assertRaises(SystemExit):
            require_worker_readiness(settings, repository)
        repository.health.side_effect = RuntimeError("credential=SECRET")
        with self.assertRaises(SystemExit) as raised:
            require_worker_readiness(settings, repository)
        self.assertNotIn("SECRET", str(raised.exception))
        repository.reset_mock()
        require_worker_readiness(SimpleNamespace(environment="development"), repository)
        repository.health.assert_not_called()

    def test_heartbeats_keep_long_job_alive(self):
        repo = FakeRepository()
        settings = self.settings.model_copy(update={"worker_lease_seconds": 0.3})
        Worker(settings, repo, self.storage, FakeProcessor(delay=0.6)).run_once()
        self.assertGreaterEqual(repo.heartbeats, 1)
        self.assertIsNotNone(repo.published)

    def test_transient_errors_retry_and_do_not_persist_provider_text(self):
        repo = FakeRepository()
        failure = httpx.ConnectError("SECRET DOCUMENT TEXT")
        Worker(self.settings, repo, self.storage, FakeProcessor(failure)).run_once()
        self.assertTrue(repo.failed["retryable"])
        self.assertNotIn("SECRET", repo.failed["message"])
        self.assertFalse(
            classify_failure(ServiceError("INVALID_DOCUMENT", "Invalid file"))[2]
        )
        wrapped = EmbeddingError("Embedding endpoint is unavailable")
        wrapped.__cause__ = failure
        self.assertTrue(classify_failure(wrapped)[2])
        invalid = EmbeddingError("Embedding response malformed")
        invalid.__cause__ = ValueError("Wrong dimension")
        self.assertFalse(classify_failure(invalid)[2])

    def test_lease_revocation_during_external_call_prevents_publication(self):
        repo = FakeRepository()
        settings = self.settings.model_copy(update={"worker_lease_seconds": 0.3})
        timer = threading.Timer(0.1, lambda: setattr(repo, "valid", False))
        timer.start()
        try:
            Worker(settings, repo, self.storage, FakeProcessor(delay=0.6)).run_once()
        finally:
            timer.cancel()
        self.assertGreaterEqual(repo.heartbeats, 1)
        self.assertIsNone(repo.published)
        self.assertIsNone(repo.failed)

    def test_cleanup_removes_objects_before_marking_complete(self):
        stored = self.storage.write(io.BytesIO(b"original"), "doc.txt")
        repo = FakeRepository(kind="cleanup")
        repo.keys = [stored.object_key]
        Worker(self.settings, repo, self.storage, FakeProcessor()).run_once()
        self.assertTrue(repo.cleaned)
        self.assertFalse((Path(self.directory.name) / stored.object_key).exists())


if __name__ == "__main__":
    unittest.main()
